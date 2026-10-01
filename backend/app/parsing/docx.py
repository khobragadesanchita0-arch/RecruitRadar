import io
from typing import Any, Dict, List
import docx
from docx.shared import RGBColor
from app.parsing.hidden_text import detect_unicode_hidden_chars, strip_hidden_unicode

def parse_docx(file_bytes: bytes) -> Dict[str, Any]:
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
    except Exception as e:
        raise ValueError(f"Corrupt or unreadable DOCX: {str(e)}")
        
    visible_text_parts = []
    hidden_text_items = []
    
    # Process paragraphs
    for p in doc.paragraphs:
        for run in p.runs:
            text = run.text
            if not text.strip():
                continue
                
            u_hidden = detect_unicode_hidden_chars(text)
            if u_hidden:
                hidden_text_items.extend(u_hidden)
                
            cleaned_text = strip_hidden_unicode(text)
            
            # Check font size
            is_tiny = False
            if run.font.size and run.font.size.pt < 2.0:
                is_tiny = True
                
            # Check white color
            is_white = False
            if run.font.color and run.font.color.rgb == RGBColor(255, 255, 255):
                is_white = True
                
            # Check hidden/vanish XML element
            is_vanish = False
            rPr = run._r.get_or_add_rPr()
            if rPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}vanish") is not None:
                is_vanish = True
                
            if is_tiny or is_white or is_vanish:
                reasons = []
                if is_tiny:
                    reasons.append("tiny_font_<2pt")
                if is_white:
                    reasons.append("white_font")
                if is_vanish:
                    reasons.append("vanish_hidden_tag")
                hidden_text_items.append({
                    "text": cleaned_text,
                    "reason": ", ".join(reasons),
                })
            else:
                visible_text_parts.append(cleaned_text)
                
    # Process tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if p.text.strip():
                        visible_text_parts.append(strip_hidden_unicode(p.text))
                        
    full_text = "\n".join(visible_text_parts).strip()
    parse_confidence = 1.0 if len(full_text) >= 100 else 0.5
    
    return {
        "visible_text": full_text,
        "hidden_text_items": hidden_text_items,
        "parse_confidence": parse_confidence,
    }
